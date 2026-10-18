from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
src=Path('output/dossier_projet/Aureo_Ruffier_Dossier_Projet_DWWM.docx')
out=src.with_name('Aureo_Ruffier_Dossier_Projet_DWWM_enrichi.docx')
with ZipFile(src) as z:data={n:z.read(n) for n in z.namelist()}
w='http://schemas.openxmlformats.org/wordprocessingml/2006/main';ns={'w':w};q=lambda s:'{'+w+'}'+s
root=E.fromstring(data['word/document.xml']);body=root.find(q('body'))
def text(e):return ''.join(e.xpath('.//w:t/text()',namespaces=ns))
def find(prefix):return next(e for e in body if e.tag==q('p') and text(e).startswith(prefix) and not e.xpath('./w:pPr/w:pStyle[starts-with(@w:val,"TOC") or starts-with(@w:val,"TM")]',namespaces=ns))
def para(t='',style=None,bold=False,size=None,keep=False):
 p=E.Element(q('p'));pp=E.SubElement(p,q('pPr'))
 if style:E.SubElement(pp,q('pStyle')).set(q('val'),style)
 if keep:E.SubElement(pp,q('keepNext'))
 r=E.SubElement(p,q('r'))
 if bold or size:
  rp=E.SubElement(r,q('rPr'))
  if bold:E.SubElement(rp,q('b'))
  if size:E.SubElement(rp,q('sz')).set(q('val'),str(size*2))
 tt=E.SubElement(r,q('t'));tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tt.text=t
 return p
def pagebreak():
 p=para();r=p.find(q('r'));E.SubElement(r,q('br')).set(q('type'),'page');return p
def insert_before(anchor,els):
 for e in els:body.insert(body.index(anchor),e)
def code(lines):
 result=[]
 for i,line in enumerate(lines):
  p=para(line,'Code',keep=i<len(lines)-1);result.append(p)
 return result
# Titles: retain numbering, bookmarks, field result and paragraph formatting.
renames={'Les actions de chaque acteur':'Acteurs et cas d’utilisation','Passer des objets aux tables':'Modèle relationnel et structure de la base','Séparer le dépôt de la publication':'Modération des avis','Ce que le projet m’a appris':'Difficultés rencontrées et solutions apportées'}
counts={}
for old,new in renames.items():
 counts[old]=0
 for p in body.findall(q('p')):
  if old not in text(p):continue
  ts=list(p.iter(q('t')));joined=''.join(t.text or '' for t in ts);start=joined.index(old);end=start+len(old);offset=0
  for t in ts:
   v=t.text or '';a=offset;b=offset+len(v);offset=b
   if b<=start or a>=end:continue
   left=v[:max(0,start-a)];right=v[max(0,end-a):] if b>end else ''
   t.text=left+(new if a<=start<b else '')+right
  counts[old]+=1
# Move SQL explanation out of migrations and pair it with the actual repository example.
start=find('Voici une requête SQL de lecture');end=find('Pour préparer les démonstrations')
for e in list(body)[body.index(start):body.index(end)]:body.remove(e)
start=find('Extrait du dépôt • src/Repository/ProductRepository.php');end=find('La lettre p représente Product')
for e in list(body)[body.index(start):body.index(end)]:body.remove(e)
D=Document(src);table=D.add_table(rows=1,cols=2);table.autofit=False
for col in table.columns:col.width=Inches(3.25)
headers=['QueryBuilder Doctrine','SQL de lecture correspondant']
php=["return $this->createQueryBuilder('p')", "    ->innerJoin('p.categories', 'c')", "    ->andWhere('p.isActive = :active')", "    ->andWhere('c.slug = :slug')", "    ->setParameter('active', true)", "    ->setParameter('slug', $slug)", "    ->orderBy('p.createdAt', 'DESC')", "    ->getQuery()", "    ->getResult();"]
sql=['SELECT p.*','FROM product AS p','INNER JOIN product_category AS pc','    ON pc.product_id = p.id','INNER JOIN category AS c','    ON c.id = pc.category_id','WHERE p.is_active = :active','  AND c.slug = :slug','ORDER BY p.created_at DESC;']
for i,h in enumerate(headers):
 cell=table.rows[0].cells[i];cell.text=h
 for run in cell.paragraphs[0].runs:run.bold=True;run.font.size=Pt(10)
 tcpr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');tcpr.append(sh)
row=table.add_row()
for i,lines in enumerate([php,sql]):
 cell=row.cells[i];cell.text='\n'.join(lines)
 for p in cell.paragraphs:
  p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
  for run in p.runs:run.font.name='Consolas';run.font.size=Pt(8)
for row in table.rows:
 pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
 for cell in row.cells:
  pp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   a=OxmlElement('w:'+side);a.set(qn('w:w'),'90');a.set(qn('w:type'),'dxa');m.append(a)
  pp.append(m)
pr=table._tbl.tblPr;b=OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
 e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
pr.append(b)
insert_before(end,[para('Une même recherche avec Doctrine et en SQL',bold=True,keep=True),para('Extrait du dépôt • src/Repository/ProductRepository.php • méthode findByCategorySlug','Source'),deepcopy(table._tbl),para('À gauche, le code utilise les propriétés des entités. À droite, la traduction pédagogique rend explicites les deux jointures SQL de la relation plusieurs à plusieurs. Les paramètres correspondent à active = true et au slug demandé. Ce SQL illustre la recherche ; il ne s’agit pas d’une capture de la requête générée par Doctrine.')])
# Add a focused service explanation as a continuation of 4.7, before its existing page break.
anchor=find('4.8 Consulter et suivre un devis');prev=anchor.getprevious()
if prev is not None and prev.xpath('.//w:br[@w:type="page"]',namespaces=ns):anchor=prev
service=[pagebreak(),para('Un service dédié au stockage des logos',bold=True,keep=True),para('J’ai regroupé la gestion des fichiers dans QuoteLogoStorage. Le contrôleur coordonne le formulaire et l’enregistrement du devis ; le service s’occupe du fichier. Symfony injecte ce service dans l’action publique et dans le constructeur du contrôleur administratif. Les deux parcours réutilisent ainsi les mêmes règles sans recopier le traitement des uploads.'),para('Le répertoire est défini dans config/services.yaml : private/quote-logos, hors du dossier public. store() vérifie la validité du fichier, sa taille maximale de 15 000 000 octets et la cohérence entre son type MIME et l’image reconnue. Seuls PNG, JPEG et WebP sont admis. Le nom généré avec random_bytes ne reprend pas le nom fourni par le visiteur. Le service demande les permissions 0700 pour le dossier et 0600 pour le fichier.'),para('Du fichier reçu à sa référence en base',bold=True,keep=True),para('Après validation du formulaire, le contrôleur transmet le fichier à store(). La méthode retourne un nom que setLogo() associe au devis. Doctrine conserve cette référence dans la base ; le contenu de l’image reste sur le disque. Le contrôleur définit ensuite la date et le statut avant persist() et flush().'),para('Extraits du dépôt • src/Controller/Front/QuoteRequestController.php • lignes 43 à 45 et 56 à 62','Source')]
service+=code(["if ($logoFile) {","    $quoteRequest->setLogo($logos->store($logoFile));","}","// Après préparation de la date et du statut","try {","    $entityManager->persist($quoteRequest);","    $entityManager->flush();","} catch (\\Throwable $exception) {","    $logos->delete($quoteRequest->getLogo());","    throw $exception;","}"])
service += [para('Si le stockage échoue, le formulaire affiche une erreur et le contrôleur renvoie une réponse 422. Si l’écriture en base échoue après le stockage, il demande la suppression du logo puis transmet l’exception. Cette compensation limite les fichiers sans devis associé ; elle ne constitue pas une transaction commune entre le disque et la base.'),para('Lors d’un remplacement dans l’administration, l’ancien logo est supprimé après la réussite de flush(). En cas d’échec, le nouveau est supprimé et l’ancienne référence rétablie dans l’objet. delete() journalise un avertissement si la suppression échoue. Pour le téléchargement, le contrôleur vérifie ROLE_ADMIN ; le service recherche un fichier autorisé. Cette séparation permet aussi de tester le stockage indépendamment du contrôleur, comme dans QuoteLogoStorageTest.')]
insert_before(anchor,service)
# Replace the short JavaScript section with an end-to-end explanation, two deliberate pages.
start=find('5.4 Ajouter des comportements');end=find('6 Sécurité et confidentialité')
a=body.index(start)+1;b=body.index(end)
for e in list(body)[a:b]:body.remove(e)
js=[para('J’utilise JavaScript pour répondre aux actions du visiteur dans le navigateur : changement de thème, agrandissement des images et aperçu d’un fichier. Pour détailler cette démarche, je prends le thème clair et sombre. Son code se trouve dans assets/scripts/theme.js, importé par assets/scripts/app.js.'),para('Repérer la commande et retrouver le choix mémorisé',bold=True,keep=True),para('Dans le template du menu, le bouton porte l’attribut data-theme-toggle. Ce repère permet de sélectionner les commandes sans dépendre de leur apparence. À l’événement DOMContentLoaded, le script lit la clé site-theme dans localStorage. Si sa valeur est dark, setTheme() active le thème sombre ; sinon, le thème clair est utilisé. La préférence est conservée dans ce navigateur pour la même origine, sans être enregistrée dans le compte utilisateur.'),para('Extrait du dépôt • assets/scripts/theme.js • lignes 31 à 43','Source')]
lines=Path('assets/scripts/theme.js').read_text().splitlines()[30:43]
js+=code(lines)
js += [para('Capturer le clic et changer l’état',bold=True,keep=True),para('querySelectorAll() récupère tous les boutons concernés. forEach() attache un écouteur click à chacun. Au clic, classList.toggle() ajoute ou retire dark-mode sur body et retourne un booléen. Je réutilise ce résultat pour choisir la valeur dark ou light, la mémoriser avec localStorage.setItem(), puis mettre à jour les commandes. Aucune requête au serveur n’est nécessaire pour ce changement.'),para('Le script modifie donc un état simple : la présence de la classe dark-mode. Le CSS utilise body.dark-mode pour redéfinir les variables de couleur, notamment celles du fond et du texte. Les composants qui utilisent ces variables changent d’apparence tout en conservant leur contenu et leur structure.'),pagebreak(),para('Synchroniser les commandes et les images du thème',bold=True,keep=True),para('updateThemeButton() actualise tous les boutons : aria-label décrit l’action suivante, title indique le mode proposé et les classes fa-moon et fa-sun changent l’icône. Lorsque le mode sombre est actif, le libellé devient « Activer le mode clair ». Le bouton reste ainsi compréhensible au-delà de son seul pictogramme.'),para('Extrait du dépôt • assets/scripts/theme.js • lignes 11 à 19','Source')]
js+=code(Path('assets/scripts/theme.js').read_text().splitlines()[10:19])
js += [para('La fonction appelle ensuite updateThemeLogos(). Les images repérées par data-theme-logo reçoivent la variante claire ou sombre de leur source. Lors du chargement d’une autre page, setTheme() applique le choix mémorisé puis appelle la même fonction de mise à jour. Le clic et l’initialisation réutilisent ainsi les mêmes règles pour les boutons et les logos.'),para('Vérifier le comportement du script',bold=True,keep=True),para('Une vérification complémentaire du fichier JavaScript a été exécutée le 17 octobre 2026 avec un DOM simulé. Elle couvre l’absence de préférence, le passage au sombre, la restauration du choix lors d’un nouveau chargement, le retour au clair et une valeur mémorisée inconnue. Les assertions ont confirmé l’état de la classe, la valeur stockée, les libellés des boutons et les sources des logos attendus.'),para('Cette vérification contrôle la logique du script ; elle ne valide pas le rendu des couleurs, le chargement réel des images ou la navigation au clavier. Dans le navigateur, ces points se contrôlent sur plusieurs pages, après rechargement et aux largeurs ordinateur et mobile. Ce contrôle complémentaire est distinct de la suite PHP présentée dans la section 7.1.'),para('Le choix technique reste adapté à cette interaction : JavaScript gère l’événement et l’état, le CSS définit l’apparence et localStorage mémorise la préférence. Cette mémorisation est locale au navigateur et ne synchronise pas le thème entre plusieurs appareils.'),pagebreak()]
insert_before(end,js)
# Preserve TOC field and images; cached page numbers require repagination in Word.
data['word/document.xml']=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
with ZipFile(out,'w',ZIP_DEFLATED) as z:
 for n,b in data.items():z.writestr(n,b)
assert all(c==2 for c in counts.values()),counts
check=Document(out)
assert len(check.inline_shapes)==len(D.inline_shapes)
assert len(root.xpath('//w:instrText[contains(text(),"TOC")]',namespaces=ns))==1
print(out.resolve());print('Renamed titles and TOC entries:',counts);print('Images preserved:',len(check.inline_shapes));print('Paragraphs:',len(check.paragraphs),'Tables:',len(check.tables))
