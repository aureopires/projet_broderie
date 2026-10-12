from pathlib import Path
import re,textwrap,json
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path.cwd(); OUT=ROOT/'output/dossier_projet';OUT.mkdir(exist_ok=True)
raw=(ROOT/'tmp/dossier/contenu.txt').read_text(encoding='utf8')
parts=re.split(r'^# ',raw,flags=re.M)[1:]
pages=[(s.split('\n',1)[0],s.split('\n',1)[1]) for s in parts]
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.78)
sec.header_distance=sec.footer_distance=Inches(.30)
styles=d.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(7)
styles['Normal'].font.size=Pt(11)
styles['Normal'].paragraph_format.line_spacing=1.06
styles['Normal'].paragraph_format.widow_control=True
styles['Title'].font.size=Pt(30)
styles['Heading 1'].font.size=Pt(20);styles['Heading 1'].paragraph_format.space_after=Pt(14)
styles['Heading 2'].font.size=Pt(12);styles['Heading 2'].paragraph_format.space_before=Pt(8)
styles['Caption'].font.size=Pt(9);styles['Caption'].font.italic=True
for style in styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for name in ['Code','Source','TocEntry']:
 if name not in styles:styles.add_style(name,1)
styles['Code'].font.name='Consolas';styles['Code'].font.size=Pt(9)
styles['Code'].paragraph_format.line_spacing=1
styles['Code'].paragraph_format.space_after=Pt(0)
styles['Source'].font.name='Calibri';styles['Source'].font.size=Pt(8)
styles['Source'].paragraph_format.space_before=Pt(6)
styles['TocEntry'].font.size=Pt(11)
styles['TocEntry'].paragraph_format.space_after=Pt(8)
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'fr-FR');styles['Normal'].element.get_or_add_rPr().append(lang)
sec.different_first_page_header_footer=True
hp=sec.header.paragraphs[0];hp.text='AUREO RUFFIER    /    DOSSIER DE PROJET DWWM'
hp.runs[0].font.size=Pt(8)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('L’Aiguille des Volcans     •     ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f)
for r in fp.runs:r.font.size=Pt(8)

def para(t,style=None):
 return d.add_paragraph(t,style)
def pic(path,caption,maxh):
 im=Image.open(ROOT/path);w,h=im.size
 width=min(6.9,float(maxh)*w/h)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
 run=p.add_run();shape=run.add_picture(str(ROOT/path),width=Inches(width))
 shape._inline.docPr.set('descr',caption)
 p=para(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def table(rows):
 n=len(rows[0]);t=d.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=([2.2,2.85,1.85] if n==3 else [2.65,1.42,1.42,1.41])
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,text in enumerate(row):
   cell=cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for tag in ['top','left','bottom','right']:
    e=OxmlElement('w:'+tag);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
   tcPr.append(margins)
   borders=OxmlElement('w:tcBorders')
   for tag in ['top','left','bottom','right']:
    e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'343434' if i==0 else ('F4F4F4' if i%2==0 else 'FFFFFF'));tcPr.append(shade)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
   r=p.add_run(text);r.font.size=Pt(9);r.bold=i==0;r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
  trPr=t.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trPr.append(no)
  if i==0:
   rep=OxmlElement('w:tblHeader');trPr.append(rep)
 para('').paragraph_format.space_after=Pt(0)
def code(spec):
 path,key,num=spec.split('|');lines=(ROOT/path).read_text(encoding='utf-8-sig').splitlines()
 start=next(i for i,s in enumerate(lines) if key in s);sel=lines[start:start+int(num)]
 nonempty=[len(s)-len(s.lstrip()) for s in sel if s.strip()];indent=min(nonempty or [0])
 para(f'Extrait du dépôt • {path} • lignes {start+1} à {start+len(sel)}','Source')
 for i,s in enumerate(sel):
  wraps=textwrap.wrap(s[indent:],width=86,replace_whitespace=False,drop_whitespace=False) or ['']
  for j,line in enumerate(wraps):
   p=para((f'{start+i+1:>3}  ' if j==0 else '     ')+line,'Code');p.paragraph_format.keep_with_next=True
 para('').paragraph_format.space_after=Pt(2)
def body(txt):
 lines=txt.splitlines();i=0;buf=[]
 def flush():
  if buf:para(' '.join(buf));buf.clear()
 while i<len(lines):
  s=lines[i]
  if s.startswith('@table:'):
   flush();rows=[s[7:].split('|')];i+=1
   while lines[i]!='@endtable':rows.append(lines[i].split('|'));i+=1
   table(rows)
  elif s.startswith('@image:'):
   flush();path,cap,h=s[7:].split('|');pic(path,cap,h)
  elif s.startswith('@code:'):flush();code(s[6:])
  elif s.startswith('@source:'):flush();para(s[8:],'Source')
  elif s.startswith('## '):flush();d.add_heading(s[3:],2)
  elif not s.strip():flush()
  else:buf.append(s)
  i+=1
 flush()

para('TITRE PROFESSIONNEL  •  RNCP37674').runs[0].font.size=Pt(11)
para('Dossier de projet','Title')
para('L’Aiguille des Volcans','Subtitle').runs[0].font.size=Pt(24)
para('Site web de broderie et de personnalisation textile').runs[0].font.size=Pt(15)
para('Développeur web et web mobile\nNiveau 5').runs[0].font.size=Pt(13)
pic('tmp/dossier/home.png','Projet Symfony et MariaDB',2.7)
para('Aureo Ruffier').runs[0].bold=True
para('L’Institut de la PME\nFormation du 4 mai au 14 août 2026\nSoutenance prévue le 29 octobre 2026')
para('Réalisation individuelle\nVersion documentaire du 11 octobre 2026\nÉtat du code étudié b406ce8')
para('Document de travail à actualiser après les correctifs et la recette finale.','Source')

for chunk in range(2):
 d.add_page_break();d.add_heading('Sommaire' if chunk==0 else 'Sommaire et figures',1)
 for idx,(title,txt) in enumerate(pages[chunk*18:(chunk+1)*18],start=chunk*18+1):
  p=para(f'{idx:02d}   {title}','TocEntry')
  p.paragraph_format.tab_stops.add_tab_stop(Inches(6.7),WD_ALIGN_PARAGRAPH.RIGHT)
  p.add_run('\t'+str(idx+3))
 if chunk==1:
  para('Annexes : A Modèle physique p. 40 ; B Maquette web p. 41 ; C Extraits du formulaire p. 42 ; D Maquette mobile p. 43.','Source')
  para('Figures 1 à 9 : architecture, parcours, wireframes, maquette mobile, accueil, rendu étroit, galerie, modèle conceptuel et séquence du devis.','Source')

for i,(title,txt) in enumerate(pages,1):
 d.add_page_break();d.add_heading(f'{i:02d}  {title}',1);body(txt)

d.add_page_break();d.add_heading('Annexe A Modèle physique des relations',1)
para('Ce schéma synthétique représente les clés et les relations de la version examinée. Les attributs détaillés figurent au chapitre 17. PK désigne une clé primaire, FK une clé étrangère et UQ une contrainte unique. Les attributs de compte et de devis contenant des données personnelles ne sont pas illustrés par des valeurs réelles.')
pic('tmp/dossier/mpd.png','Figure A1 Schéma physique simplifié reconstitué à partir du mappage Doctrine',4.4)
para('Les relations de product_category utilisent une clé primaire composée. review.user_id est obligatoire, tandis que review.product_id et quote_request.user_id peuvent être NULL. Le lien produit vers avis prévoit ON DELETE CASCADE. La table technique messenger_messages est volontairement séparée du modèle métier.')
para('Le contrôle doctrine:schema:validate a confirmé le mappage et la synchronisation de la base locale. Il ne garantit pas l’absence de perte dans toutes les migrations historiques.','Source')
d.add_page_break();d.add_heading('Annexe B Maquette web du devis',1)
para('Support graphique déjà fourni dans le projet, conservé pour la comparaison avec le formulaire et son adaptation mobile. Cette figure ne constitue pas une capture de la production actuelle.')
pic('tmp/dossier/maquette-web.png','Figure B1 Maquette web du devis issue du support SVG fourni',6.65)
d.add_page_break();d.add_heading('Annexe C Extraits du formulaire de devis',1)
para('Les extraits ci-dessous montrent la séparation entre le fichier temporaire non mappé et la quantité persistante. Ils sont reproduits depuis le dépôt et doivent être lus avec les résultats des essais du chapitre 29.')
code("src/Form/QuoteRequestType.php|->add('logoFile'|14")
code("src/Form/QuoteRequestType.php|->add('quantity'|16")
para('Le type File vérifie la taille et une liste de MIME. Il ne rend pas automatiquement le SVG inoffensif. GreaterThanOrEqual contrôle la quantité côté serveur ; required ne remplace pas une contrainte NotBlank sur les champs textuels. Les corrections proposées doivent rester distinctes de ces extraits de la version examinée.')
d.add_page_break();d.add_heading('Annexe D Maquette mobile du devis',1)
para('Adaptation mobile fournie dans les supports graphiques du projet. Les deux parties se lisent de gauche à droite. Elles présentent une intention de conception et ne remplacent pas une recette sur téléphone.')
pic('tmp/dossier/maquette-mobile.png','Figure D1 Maquette mobile complète du formulaire en deux parties',7.0)

d.core_properties.title='Dossier de projet L’Aiguille des Volcans'
d.core_properties.author='Aureo Ruffier'
d.core_properties.subject='Titre professionnel DWWM RNCP37674'
d.save(OUT/'Aureo_Ruffier_Dossier_Projet_DWWM.docx')
print(json.dumps({'body_sections':len(pages),'expected_pages':len(pages)+7,'words':len(raw.split()),'output':str(OUT)},ensure_ascii=True))

