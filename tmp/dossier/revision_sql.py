from pathlib import Path
from copy import deepcopy
import shutil
from docx import Document

root = Path.cwd()
source = root / 'output/dossier_projet/Aureo_Ruffier_Dossier_Projet_DWWM.docx'
backup = root / 'tmp/dossier/avant_revision_sql.docx'
if not backup.exists():
    shutil.copy2(source, backup)
doc = Document(source)

def replace_text(paragraph, text):
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if properties is not None:
        run._r.insert(0, properties)

changes = 0
for paragraph in doc.paragraphs:
    text = paragraph.text
    if text.startswith('La présence de JSON dans une colonne MariaDB'):
        replace_text(paragraph, 'La matrice présente les compétences mobilisées dans le périmètre du projet, avec des intitulés abrégés. L’accès aux données est illustré par les repositories Doctrine et les requêtes exécutées sur MariaDB. Les exemples relient la sélection, l’enregistrement et la modification des données aux fonctionnalités du catalogue, des devis et des avis.')
        changes += 1
    elif 'Aucun composant d’accès à une base NoSQL' in text:
        replace_text(paragraph, text.replace(' Aucun composant d’accès à une base NoSQL n’a été identifié dans le périmètre développé.', ''))
        changes += 1
    elif 'Une fonctionnalité NoSQL éventuelle' in text:
        replace_text(paragraph, text.replace(' Une fonctionnalité NoSQL éventuelle doit répondre à un besoin réel et être présentée avec ses propres preuves.', ''))
        changes += 1

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.text == 'CP6 Accéder aux données SQL et NoSQL':
                    replace_text(paragraph, 'CP6 Développer l’accès aux données')
                    changes += 1
                elif paragraph.text == 'SQL démontré ; NoSQL non réalisé':
                    replace_text(paragraph, 'Accès SQL avec MariaDB et Doctrine')
                    changes += 1

text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
assert 'NoSQL' not in text
assert changes == 5, changes
doc.save(source)

content_path = root / 'tmp/dossier/contenu.txt'
content = content_path.read_text(encoding='utf-8')
content = content.replace('CP6 Accéder aux données SQL et NoSQL', 'CP6 Développer l’accès aux données').replace('SQL démontré ; NoSQL non réalisé', 'Accès SQL avec MariaDB et Doctrine')
lines = content.splitlines()
for index, line in enumerate(lines):
    if line.startswith('La présence de JSON dans une colonne MariaDB'):
        lines[index] = 'La matrice présente les compétences mobilisées dans le périmètre du projet, avec des intitulés abrégés. L’accès aux données est illustré par les repositories Doctrine et les requêtes exécutées sur MariaDB. Les exemples relient la sélection, l’enregistrement et la modification des données aux fonctionnalités du catalogue, des devis et des avis.'
    else:
        lines[index] = line.replace(' Aucun composant d’accès à une base NoSQL n’a été identifié dans le périmètre développé.', '').replace(' Une fonctionnalité NoSQL éventuelle doit répondre à un besoin réel et être présentée avec ses propres preuves.', '')
assert 'NoSQL' not in '\n'.join(lines)
content_path.write_text('\n'.join(lines) + '\n', encoding='utf-8')
print('Dossier updated: 5 text changes; no NoSQL mentions; formatting preserved.')
