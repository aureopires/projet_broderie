from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from PIL import Image,ImageDraw,ImageFont
from zipfile import ZipFile
from lxml import etree as E
src=Path('output/dossier_projet/Aureo_Ruffier_Dossier_Projet_DWWM.docx')
out=src.with_name('Aureo_Ruffier_Dossier_Projet_DWWM_tests_documentes.docx')
report=Path('output/tests/phpunit-2026-10-17_19-07-09.txt')
log=report.read_text(encoding='utf-8-sig');assert 'OK (63 tests, 229 assertions)' in log
lines=log.splitlines()
selected=[next(x for x in lines if x.startswith('PHPUnit')),next(x for x in lines if x.startswith('Runtime:')),next(x for x in lines if x.startswith('Configuration:')),'',next(x for x in lines if '63 / 63' in x),next(x for x in lines if x.startswith('Time:')),'',next(x.strip() for x in lines if 'data set "quantité zéro"' in x),'',next(x for x in lines if x.startswith('OK ('))]
font=ImageFont.truetype('C:/Windows/Fonts/consola.ttf',24)
im=Image.new('RGB',(1480,410),'white');draw=ImageDraw.Draw(im)
for i,t in enumerate(selected):draw.text((24,18+36*i),t.lstrip('✔ ').lstrip('� '),font=font,fill='#111111')
img=Path('output/tests/extrait-phpunit-2026-10-17.png');im.save(img)
d=Document(src)
for p in d.paragraphs:
 if p.text.startswith('Je vérifie les parcours dans le navigateur'):
  p.text='Pendant le développement, j’ai vérifié manuellement les principaux parcours du site, notamment la navigation, les formulaires et l’affichage sur différentes tailles d’écran. En complément, une suite de tests automatisés a été exécutée sur la version locale. Ses résultats ne constituent pas un décompte de mes essais manuels.'
 elif p.text.startswith('La suite du projet a été exécutée'):
  p.text='La suite PHPUnit a été exécutée de nouveau le 17 octobre 2026 à 19 h 07, heure de Paris : 63 tests et 229 assertions ont réussi. Le rapport complet est conservé dans output/tests/phpunit-2026-10-17_19-07-09.txt. Les contrôles Doctrine et des 54 fichiers Twig mentionnés dans la version précédente du dossier se rapportent à la vérification du 13 octobre ; ils ne font pas partie de cette nouvelle exécution de PHPUnit.'
 elif p.text.startswith('J’ai donc utilisé les deux approches.'):
  p.text='Les deux approches se complètent : mes essais manuels portent sur les écrans et leur utilisation, tandis que les tests automatisés vérifient les règles et les scénarios définis dans le code. Leur réussite ne garantit pas l’absence de toute anomalie et ne remplace pas une vérification visuelle.'
 elif p.text.startswith('Ces résultats proviennent de la suite locale exécutée le 13 octobre'):
  p.text=p.text.replace('13 octobre','17 octobre')
anchor=next(p for p in d.paragraphs if p.text.startswith('7.2 ') and p.style.style_id=='Titre2')
# Retain original break between the two numbered sections after the new continuation page.
prev=anchor._p.getprevious()
if prev is not None and prev.xpath('.//w:br[@w:type="page"]'):insert_anchor=prev
else:insert_anchor=anchor._p
new=[]
def add(t='',style=None,bold=False):
 p=d.add_paragraph(t,style);new.append(p._p)
 if bold:
  for r in p.runs:r.bold=True
  p.paragraph_format.keep_with_next=True
 return p
add().add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
add('Exécution documentée et exemple de validation',bold=True)
add('La commande ci-dessous exécute la suite existante dans le service PHP de Docker. L’option --testdox affiche les scénarios ; --colors=never produit un rapport sans codes de couleur. Le fichier texte complet conserve la sortie de cette exécution.')
p=add('docker compose exec -T php php bin/phpunit --testdox --colors=never','Code')
for r in p.runs:r.font.size=Pt(9)
p=add();p.add_run().add_picture(str(img),width=Inches(6.65))
add('Extrait de la sortie réelle du 17 octobre 2026 à 19 h 07, reproduit depuis le rapport texte. Les lignes intermédiaires sont omises.','Caption')
add('Règle vérifiée et préparation',bold=True)
add('Une demande de devis doit porter sur au moins un article. QuoteRequestTypeTest utilise des données fictives valides, dont une quantité de douze polos. Le fournisseur invalidData() transmet ensuite le cas « quantité zéro » : field vaut quantity et value vaut 0. Le test remplace uniquement cette valeur avant de soumettre le formulaire Symfony.')
add('Extraits de tests/Form/QuoteRequestTypeTest.php • cas de données et assertions','Source')
for line in ["yield 'quantité zéro' => ['quantity', 0];","",'$form->submit(array_replace(self::VALID_DATA, [$field => $value]));','self::assertFalse($form->isValid());','self::assertGreaterThan(0, count($form->get($field)->getErrors(true)));']:
 p=add(line,'Code')
 for r in p.runs:r.font.size=Pt(9)
 p.paragraph_format.keep_with_next=True
add('Action et résultats attendus',bold=True)
add('submit() déclenche le traitement et la validation du formulaire. La première assertion exige que isValid() soit faux. La seconde exige au moins une erreur sur le champ quantity. Le scénario « quantité zéro » a réussi : ces deux conditions sont vérifiées. Le même test est réutilisé pour d’autres données invalides, ce qui explique plusieurs scénarios à partir d’une seule méthode.')
add('Portée du test',bold=True)
add('Ce test vérifie la validation côté serveur, sans navigateur visuel. Le CSRF est désactivé uniquement dans ce test isolé pour cibler les contraintes des champs. Il ne vérifie pas une écriture en base : le scénario fonctionnel distinct testInvalidQuoteDoesNotPersist() contrôle qu’une demande avec un e-mail invalide et un prénom vide n’est pas enregistrée.')
for el in new:insert_anchor.addprevious(el)
d.save(out)
check=Document(out)
assert len(check.inline_shapes)==len(Document(src).inline_shapes)+1
with ZipFile(src) as z:original=E.fromstring(z.read('word/document.xml'))
with ZipFile(out) as z:
 assert z.testzip() is None
 result=E.fromstring(z.read('word/document.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
assert original.xpath('//w:instrText/text()',namespaces=ns)==result.xpath('//w:instrText/text()',namespaces=ns)
assert len(original.xpath('//w:bookmarkStart',namespaces=ns))==len(result.xpath('//w:bookmarkStart',namespaces=ns))
print(out.resolve());print('Verified: original pictures and TOC fields retained, one evidence image added, report totals confirmed. No PDF created.')
